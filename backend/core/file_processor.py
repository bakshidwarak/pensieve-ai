"""File processing utilities for extracting text from various file types."""

import os
import io
from typing import Optional, Tuple
from PIL import Image
import pytesseract
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation


class FileProcessor:
    """Process various file types and extract text content."""

    @staticmethod
    def process_image(file_path: str) -> str:
        """
        Extract text from image using OCR.

        Args:
            file_path: Path to image file

        Returns:
            Extracted text
        """
        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            return text.strip()
        except Exception as e:
            return f"Error extracting text from image: {str(e)}"

    @staticmethod
    def process_pdf(file_path: str) -> str:
        """
        Extract text from PDF.

        Args:
            file_path: Path to PDF file

        Returns:
            Extracted text
        """
        try:
            reader = PdfReader(file_path)
            text_parts = []

            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

            return "\n\n".join(text_parts).strip()
        except Exception as e:
            return f"Error extracting text from PDF: {str(e)}"

    @staticmethod
    def process_text(file_path: str) -> str:
        """
        Read text from plain text file.

        Args:
            file_path: Path to text file

        Returns:
            File content
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read().strip()
        except Exception as e:
            return f"Error reading text file: {str(e)}"

    @staticmethod
    def process_docx(file_path: str) -> str:
        """
        Extract text from DOCX file.

        Args:
            file_path: Path to DOCX file

        Returns:
            Extracted text
        """
        try:
            doc = Document(file_path)
            text_parts = []

            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)

            return "\n\n".join(text_parts).strip()
        except Exception as e:
            return f"Error extracting text from DOCX: {str(e)}"

    @staticmethod
    def process_pptx(file_path: str) -> str:
        """
        Extract text from PPTX file.

        Args:
            file_path: Path to PPTX file

        Returns:
            Extracted text
        """
        try:
            prs = Presentation(file_path)
            text_parts = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text_parts.append(shape.text)

            return "\n\n".join(text_parts).strip()
        except Exception as e:
            return f"Error extracting text from PPTX: {str(e)}"

    @classmethod
    def process_file(cls, file_path: str, file_type: str) -> Tuple[str, str]:
        """
        Process file based on type and extract text.

        Args:
            file_path: Path to file
            file_type: Type of file (image, pdf, text, docx, pptx)

        Returns:
            Tuple of (extracted_text, content_type)
        """
        processors = {
            "image": cls.process_image,
            "pdf": cls.process_pdf,
            "text": cls.process_text,
            "docx": cls.process_docx,
            "pptx": cls.process_pptx,
        }

        processor = processors.get(file_type)
        if not processor:
            return f"Unsupported file type: {file_type}", "error"

        text = processor(file_path)
        return text, file_type

    @staticmethod
    def get_file_type(filename: str) -> Optional[str]:
        """
        Determine file type from filename extension.

        Args:
            filename: Name of file

        Returns:
            File type category or None
        """
        ext = os.path.splitext(filename)[1].lower()

        type_mapping = {
            ".png": "image",
            ".jpg": "image",
            ".jpeg": "image",
            ".gif": "image",
            ".bmp": "image",
            ".pdf": "pdf",
            ".txt": "text",
            ".md": "text",
            ".docx": "docx",
            ".doc": "docx",
            ".pptx": "pptx",
            ".ppt": "pptx",
        }

        return type_mapping.get(ext)
